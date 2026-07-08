#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
 OUTIL DE RAPPROCHEMENT AUTOMATIQUE - PRODUCTION ASSURANCES (VERSION CORRIGÉE)
================================================================================
Corrections principales :
  - Gestion robuste des DataFrames vides
  - Vérification de l'existence des colonnes avant accès
  - Meilleure gestion des erreurs de conversion
  - Ajout de valeurs par défaut partout
================================================================================
"""

import os
import re
import sys
import io
import glob
import configparser
import unicodedata
import traceback
from datetime import datetime, timedelta
from difflib import SequenceMatcher
from collections import defaultdict

try:
    import pandas as pd
except ImportError:
    print("ERREUR : le module 'pandas' n'est pas installé.")
    print("Ouvrez une invite de commande et tapez : pip install pandas openpyxl pdfplumber")
    sys.exit(1)

try:
    import pdfplumber
except ImportError:
    print("ERREUR : le module 'pdfplumber' n'est pas installé.")
    print("Ouvrez une invite de commande et tapez : pip install pandas openpyxl pdfplumber")
    sys.exit(1)

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# Import de la configuration des sources
try:
    from config_sources import SourcesData
    SOURCES_DISPONIBLES = True
except ImportError:
    SOURCES_DISPONIBLES = False


# ================================================================================
# JOURNAL DE DIAGNOSTIC
# ================================================================================
JOURNAL = []


def log(message):
    horodatage = datetime.now().strftime("%H:%M:%S")
    ligne = f"[{horodatage}] {message}"
    JOURNAL.append(ligne)
    print(ligne)


def ecrire_journal(dossier_sortie):
    chemin = os.path.join(dossier_sortie, "journal_diagnostic.txt")
    try:
        with open(chemin, "w", encoding="utf-8") as f:
            f.write("\n".join(JOURNAL))
        log(f"Journal de diagnostic écrit : {chemin}")
    except Exception as e:
        print(f"Impossible d'écrire le journal : {e}")


# ================================================================================
# OUTILS DE NORMALISATION / COMPARAISON DE TEXTE
# ================================================================================

def normaliser_texte(valeur):
    if valeur is None:
        return ""
    texte = str(valeur).strip()
    texte = unicodedata.normalize("NFKD", texte).encode("ASCII", "ignore").decode("ASCII")
    texte = re.sub(r"\s+", " ", texte).strip().upper()
    return texte


def normaliser_cle(valeur):
    """Normalise un numéro (police, attestation, quittance)."""
    if valeur is None:
        return ""
    texte = str(valeur).strip().upper()
    texte = re.sub(r"[^A-Z0-9]", "", texte)
    return texte


def collapse_doubled(texte):
    """'RReessttee' -> 'Reste' (le PDF imprime certains mots en double épaisseur)."""
    return re.sub(r"(.)\1", r"\1", texte)


def cle_attestation_depuis_nom_fichier(nom_fichier):
    """Extrait le N° d'attestation depuis le nom du fichier scanné."""
    base = os.path.splitext(os.path.basename(nom_fichier))[0]
    base = re.sub(r"\(\d+\)\s*$", "", base).strip()

    # Cas 1 : "Lettre" puis "numéro"
    m = re.match(r"^([A-Za-z])[\s_\-]+(\d{6,12})", base)
    if m:
        return normaliser_cle(m.group(1) + m.group(2))

    # Cas 2 : numéro de quittance type "2026-01841"
    m = re.match(r"^(\d{4}-\d{2,6})", base)
    if m:
        return normaliser_cle(m.group(1))

    # Cas 3 (repli) : on garde le nom tel quel
    return normaliser_cle(base)


def similarite(a, b):
    a, b = normaliser_texte(a), normaliser_texte(b)
    if not a or not b:
        return 0
    return SequenceMatcher(None, a, b).ratio() * 100


def parser_montant(valeur):
    if valeur is None or valeur == "":
        return None
    if isinstance(valeur, (int, float)):
        return float(valeur)
    texte = str(valeur).strip()
    texte = texte.replace(" ", "").replace(" ", "").replace("\xa0", "")
    texte = texte.replace(",", ".")
    texte = re.sub(r"[^0-9.\-]", "", texte)
    if texte in ("", "-", "."):
        return None
    try:
        return float(texte)
    except ValueError:
        return None


def parser_date(valeur):
    if valeur is None or valeur == "":
        return None
    if isinstance(valeur, datetime):
        return valeur
    texte = str(valeur).strip()
    if re.search(r"_", texte) or texte.count("*") >= 2:
        return None
    formats = [
        "%d/%m/%y", "%d/%m/%Y", "%d-%m-%y", "%d-%m-%Y",
        "%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(texte, fmt)
        except ValueError:
            continue
    try:
        serie = float(texte)
        if 20000 < serie < 60000:
            return datetime(1899, 12, 30) + timedelta(days=serie)
    except ValueError:
        pass
    return None


# ================================================================================
# CONFIGURATION
# ================================================================================

def charger_config(chemin="config.ini"):
    config = configparser.ConfigParser()
    valeurs_defaut = {
        "Chemins": {"drive_path": ".", "scans_path": ".", "output_path": "./Rapports_Generes"},
        "Fichiers": {
            "pattern_encaissements": "Etat des encaissements*.pdf;Etat_des_encaissements*.pdf;Rapport encaissement*.pdf",
            "pattern_matu": "RAPPORT MATU*.xlsx",
            "pattern_sanlam": "RAPPORT SANLAM*.xlsx",
            "pattern_wafa": "RAPPORT WAFA*.xlsx",
            "pattern_maroc_assistance": "RAPPORT MAROC ASSISTANCE*.xlsx",
        },
        "Tolerances": {"tolerance_prime": "2.0", "tolerance_jours": "1", "seuil_similarite_nom": "80"},
        "Controle": {"date_controle": "auto", "filtrer_scans_par_date": "non"},
    }
    if not os.path.exists(chemin):
        log(f"ATTENTION : fichier {chemin} introuvable, valeurs par défaut utilisées.")
        for section, valeurs in valeurs_defaut.items():
            config[section] = valeurs
        return config

    config.read(chemin, encoding="utf-8")
    for section, valeurs in valeurs_defaut.items():
        if section not in config:
            config[section] = {}
        for cle, val in valeurs.items():
            if cle not in config[section]:
                config[section][cle] = val
    return config


# ================================================================================
# RECHERCHE AUTOMATIQUE DES FICHIERS
# ================================================================================

def _motif_vers_regex(motif):
    motif_echappe = re.escape(motif)
    motif_echappe = motif_echappe.replace(r"\*", ".*")
    motif_echappe = motif_echappe.replace(r"\ ", r"[ _\-]*")
    return motif_echappe


def _extraire_date_de_nom_fichier(nom):
    base = os.path.splitext(nom)[0]
    motifs = [
        r"(\d{1,2})[_\-\.](\d{1,2})[_\-\.](\d{4})",
        r"(\d{1,2})[_\-\.](\d{1,2})[_\-\.](\d{2})(?!\d)",
        r"(\d{2})(\d{2})(\d{4})",
    ]
    for motif in motifs:
        m = re.search(motif, base)
        if not m:
            continue
        j, mo, a = (int(x) for x in m.groups())
        if a < 100:
            a += 2000
        try:
            return datetime(a, mo, j).date()
        except ValueError:
            continue
    return None


def extraire_date_scan(nom_fichier):
    base = os.path.splitext(os.path.basename(nom_fichier))[0]
    base = re.sub(r"\(\d+\)\s*$", "", base).strip()
    segments = re.split(r"[_\s]+", base)
    if len(segments) >= 3 and re.fullmatch(r"\d{6}", segments[2]):
        j, mo, a = segments[2][:2], segments[2][2:4], segments[2][4:6]
        try:
            return datetime(2000 + int(a), int(mo), int(j)).date()
        except ValueError:
            return None
    return None


def _resoudre_date_controle(valeur):
    valeur_brute = (valeur or "auto").strip()
    if valeur_brute.lower() in ("auto", "", "aujourd'hui", "today"):
        return datetime.now().date()
    d = parser_date(valeur_brute)
    if d:
        return d.date()
    log(f"ATTENTION : date_controle '{valeur_brute}' non reconnue, utilisation d'aujourd'hui.")
    return datetime.now().date()


def demander_date_controle(valeur_config):
    valeur_defaut = _resoudre_date_controle(valeur_config)
    suggestion = valeur_defaut.strftime("%d/%m/%Y")

    print()
    print("=" * 70)
    try:
        reponse = input(f"Quel jour voulez-vous contrôler ? (JJ/MM/AAAA) [Entrée = {suggestion}] : ").strip()
    except EOFError:
        reponse = ""
    print("=" * 70)

    if not reponse:
        date_choisie = valeur_defaut
    else:
        d = parser_date(reponse)
        if d:
            date_choisie = d.date()
        else:
            print(f"Date '{reponse}' non reconnue, utilisation de {suggestion}.")
            date_choisie = valeur_defaut

    log(f"Jour à contrôler : {date_choisie.strftime('%d/%m/%Y')}")
    return date_choisie


def trouver_fichier_pour_date(dossier, motifs, date_cible):
    if not os.path.isdir(dossier):
        log(f"ATTENTION : dossier introuvable -> {dossier}")
        return None, False

    liste_motifs = [m.strip() for m in motifs.split(";") if m.strip()]
    candidats = []
    try:
        noms = os.listdir(dossier)
    except Exception as e:
        log(f"ERREUR en parcourant {dossier} : {e}")
        return None, False

    for motif in liste_motifs:
        regex = _motif_vers_regex(motif)
        for nom in noms:
            chemin = os.path.join(dossier, nom)
            if not os.path.isfile(chemin):
                continue
            if re.match(regex, nom, re.IGNORECASE):
                candidats.append(chemin)

    candidats = list(set(candidats))
    if not candidats:
        log(f"Aucun fichier ne correspond à {liste_motifs} dans {dossier}")
        return None, False

    # Correspondance par la date dans le NOM du fichier
    candidats_avec_date = [(c, _extraire_date_de_nom_fichier(os.path.basename(c))) for c in candidats]
    correspondances = [c for c, d in candidats_avec_date if d == date_cible]
    if len(correspondances) == 1:
        log(f"Trouvé pour {liste_motifs}, daté du {date_cible.strftime('%d/%m/%Y')} -> {os.path.basename(correspondances[0])}")
        return correspondances[0], True
    if len(correspondances) > 1:
        plus_recent = max(correspondances, key=os.path.getmtime)
        log(f"Plusieurs fichiers datés du {date_cible.strftime('%d/%m/%Y')}, utilisation du plus récent.")
        return plus_recent, True

    # Repli pour les PDF
    candidats_pdf = [c for c in candidats if c.lower().endswith(".pdf")]
    if candidats_pdf:
        for c in candidats_pdf:
            try:
                with pdfplumber.open(c) as pdf:
                    texte = pdf.pages[0].extract_text() or ""
                m = re.search(r"DU\s+(\d{2}/\d{2}/\d{2,4})\s+AU\s+(\d{2}/\d{2}/\d{2,4})", texte)
                if m:
                    d = parser_date(m.group(1))
                    if d and d.date() == date_cible:
                        log(f"Trouvé pour {liste_motifs} (date lue dans le PDF) -> {os.path.basename(c)}")
                        return c, True
            except Exception:
                continue

    # Repli final
    plus_recent = max(candidats, key=os.path.getmtime)
    log(f"** ATTENTION ** : aucun fichier daté du {date_cible.strftime('%d/%m/%Y')} trouvé. "
        f"Utilisation du plus récent -> {os.path.basename(plus_recent)}")
    return plus_recent, False


def lister_fichiers_scans(dossier, date_cible=None, filtrer_par_date=True):
    scans_jour, scans_tous = {}, {}
    if not os.path.isdir(dossier):
        log(f"ATTENTION : dossier de scans introuvable -> {dossier}")
        return scans_jour, scans_tous

    nb_dates_reconnues = 0
    try:
        noms = os.listdir(dossier)
    except Exception as e:
        log(f"ERREUR en parcourant le dossier de scans : {e}")
        return scans_jour, scans_tous

    for nom in noms:
        if not nom.lower().endswith(".pdf"):
            continue
        cle = cle_attestation_depuis_nom_fichier(nom)
        if not cle:
            continue
        scans_tous[cle] = nom
        date_scan = extraire_date_scan(nom)
        if date_scan is not None:
            nb_dates_reconnues += 1
            if date_cible is not None and date_scan == date_cible:
                scans_jour[cle] = nom

    log(f"{len(scans_tous)} fichier(s) scanné(s) au total ({nb_dates_reconnues} avec date reconnue).")
    if filtrer_par_date and date_cible is not None:
        log(f"{len(scans_jour)} scan(s) daté(s) du {date_cible.strftime('%d/%m/%Y')}.")
    return scans_jour, scans_tous


# ================================================================================
# LECTURE DU PDF "ÉTAT D'ENCAISSEMENTS" (ANALYSIS) - méthode calibrée
# ================================================================================

def _cluster_lignes(mots, tolerance=1.6):
    mots = sorted(mots, key=lambda w: w["top"])
    clusters = []
    courant = []
    top_ref = None
    for w in mots:
        if top_ref is None or abs(w["top"] - top_ref) <= tolerance:
            courant.append(w)
            top_ref = courant[0]["top"]
        else:
            clusters.append(courant)
            courant = [w]
            top_ref = w["top"]
    if courant:
        clusters.append(courant)
    return clusters


def _construire_limites_colonnes(ancres):
    items = sorted([(k, v) for k, v in ancres.items() if v is not None], key=lambda kv: kv[1])
    bornes = []
    for i, (nom, x0) in enumerate(items):
        gauche = -1e9 if i == 0 else (items[i - 1][1] + x0) / 2
        droite = 1e9 if i == len(items) - 1 else (x0 + items[i + 1][1]) / 2
        bornes.append((nom, gauche, droite))
    return bornes


def _colonne_pour_x(x0, bornes):
    for nom, gauche, droite in bornes:
        if gauche <= x0 < droite:
            return nom
    return bornes[-1][0] if bornes else "inconnu"


def _est_sous_ligne_espacee(cl, textes):
    """
    Détecte une sous-quittance imprimée caractère par caractère.
    Format: '- QUITTANCE - POLICE DATE MONTANT' avec beaucoup de tokens
    d'un seul caractère (chèque groupé : MOUNTED CAR, ERRAFIDAYNE CAR...).
    """
    if not textes or textes[0] != "-":
        return False
    mono = sum(1 for t in textes if len(t) == 1)
    return len(textes) >= 8 and mono >= 0.5 * len(textes)


def _reconstruire_champs_espaces(cl, seuil_gap=15):
    """
    Regroupe les tokens espacés en champs d'après les écarts horizontaux (x0).
    Les caractères d'un même champ sont collés ; un grand écart sépare les champs.
    Retourne la liste des champs (les séparateurs '-' isolés sont retirés).
    """
    champs = []
    courant = []
    dernier_x1 = None
    for w in sorted(cl, key=lambda w: w["x0"]):
        if dernier_x1 is not None and (w["x0"] - dernier_x1) > seuil_gap:
            if courant:
                champs.append("".join(courant))
                courant = []
        courant.append(w["text"])
        dernier_x1 = w["x1"]
    if courant:
        champs.append("".join(courant))
    # Retirer les tirets séparateurs isolés
    return [c for c in champs if c not in ("-", "")]


def extraire_etat_encaissements(chemin_pdf):
    lignes_principales = []
    lignes_rappel = []
    totaux_pdf = []
    bornes = None
    en_tete_detecte = False
    operateur_courant = ""
    periode_controlee = None

    try:
        with pdfplumber.open(chemin_pdf) as pdf:
            for num_page, page in enumerate(pdf.pages, start=1):
                mots = page.extract_words(use_text_flow=False, keep_blank_chars=False)
                if not mots:
                    continue
                clusters = _cluster_lignes(mots, tolerance=1.6)
                mode_rappel = False

                for cl in clusters:
                    cl = sorted(cl, key=lambda w: w["x0"])
                    textes = [w["text"] for w in cl]
                    ligne_jointe = " ".join(textes)
                    ligne_stripee = ligne_jointe.strip()

                    # Détection de l'en-tête
                    if "Recu" in textes and "Attestation" in textes and "Police" in textes:
                        def x0_de(jeton):
                            for w in cl:
                                if w["text"] == jeton:
                                    return w["x0"]
                            return None

                        def x0_compresse(cible):
                            for w in cl:
                                if collapse_doubled(w["text"]).lower() == cible:
                                    return w["x0"]
                            return None

                        ancres = {
                            "recu": x0_de("Recu"), "quittance": x0_de("Quitance"),
                            "attestation": x0_de("Attestation"), "assure": x0_de("Assuré"),
                            "police": x0_de("Police"), "prime": x0_de("Prime"),
                            "reference": x0_de("Reference"), "especes": x0_de("Especes"),
                            "dt": x0_de("DT"), "cheque": x0_de("Cheque"),
                            "banque": x0_de("Banque"), "operat": x0_de("Operat"),
                            "reste": x0_compresse("reste"),
                            # Colonne "Classmnt" (mise en page 30/06) : sans ancre, ses valeurs
                            # (ex 260630-006) tombaient dans "Reste" et le rendaient illisible.
                            "classmnt": x0_de("Classmnt"),
                        }
                        if ancres["police"] is not None and ancres["prime"] is not None:
                            ancres["date_effet"] = (ancres["police"] + ancres["prime"]) / 2 + 5
                        bornes = _construire_limites_colonnes(ancres)
                        en_tete_detecte = True
                        continue

                    # Lignes à ignorer
                    if ligne_stripee in ("Date", "effet"):
                        continue
                    if re.fullmatch(r"([A-Z]\s){2,}[A-Z]?", ligne_stripee):
                        continue
                    if "encaissements" in ligne_jointe.lower() and "primes" in ligne_jointe.lower():
                        continue
                    if ligne_stripee.startswith("DU "):
                        m = re.match(r"DU\s+(\d{2}/\d{2}/\d{2,4})\s+AU\s+(\d{2}/\d{2}/\d{2,4})", ligne_stripee)
                        if m and periode_controlee is None:
                            periode_controlee = (m.group(1), m.group(2))
                        continue
                    if ligne_stripee.startswith("Visa") or re.fullmatch(r"\d+\s*/\s*\d+", ligne_stripee):
                        continue

                    # Sous-tableau "Quittances encaissée"
                    if ligne_stripee == "Quittances encaissée":
                        mode_rappel = True
                        continue
                    if ligne_stripee.startswith("Police Quittance Date Prime"):
                        continue

                    # Totaux
                    if ligne_stripee.startswith("Sous totaux") or ligne_stripee.startswith("Totaux"):
                        totaux_pdf.append(ligne_stripee)
                        continue

                    # Étiquette d'opérateur
                    if len(textes) == 1 and textes[0].isalpha() and textes[0].isupper() and cl[0]["x0"] < 30:
                        operateur_courant = textes[0]
                        continue

                    # Lettres isolées
                    if all(len(t) <= 2 for t in textes) and all(t.isalpha() for t in textes) and cl[0]["x0"] > 790:
                        continue

                    # Sous-quittance imprimée caractère par caractère (chèque groupé)
                    # -> traiter comme quittance de rappel, NE PAS compter dans le total.
                    if _est_sous_ligne_espacee(cl, textes):
                        champs = _reconstruire_champs_espaces(cl)
                        if len(champs) >= 2:
                            quittance = champs[0]
                            police = champs[1]  # = attestation pour MAROC ASSISTANCE
                            date_txt = champs[2] if len(champs) >= 3 else ""
                            montant_txt = champs[3] if len(champs) >= 4 else ""
                            lignes_rappel.append({
                                "police": police,
                                "assure": "",
                                "quittance": quittance,
                                "date_effet": parser_date(date_txt),
                                "prime": parser_montant(montant_txt) if montant_txt else None,
                                "montant_encaisse": parser_montant(montant_txt) if montant_txt else None,
                            })
                        continue

                    if not bornes:
                        continue

                    premier_jeton = textes[0] if textes else ""
                    est_ligne_principale = bool(re.match(r"^\d{4}-\d{2,6}$", premier_jeton))

                    if mode_rappel and not est_ligne_principale:
                        if len(textes) >= 5:
                            lignes_rappel.append({
                                "police": textes[0],
                                "assure": "",  # rempli ensuite via le rapport compagnie (par police)
                                "quittance": textes[1],
                                "date_effet": parser_date(textes[2]),
                                "prime": parser_montant(textes[3]),
                                "montant_encaisse": parser_montant(textes[4]),
                            })
                        continue
                    else:
                        mode_rappel = False

                    colonnes = defaultdict(list)
                    for w in cl:
                        nom_col = _colonne_pour_x(w["x0"], bornes)
                        colonnes[nom_col].append(w["text"])
                    ligne = {k: " ".join(v) for k, v in colonnes.items()}

                    lignes_principales.append({
                        "quittance": ligne.get("recu", ""),
                        "quittance_compagnie": ligne.get("quittance", ""),
                        "attestation": ligne.get("attestation", ""),
                        "assure": ligne.get("assure", ""),
                        "police": ligne.get("police", ""),
                        "date_effet": parser_date(ligne.get("date_effet")),
                        "prime": parser_montant(ligne.get("prime")),
                        "reference_banque": ligne.get("reference", ""),
                        "especes": parser_montant(ligne.get("especes")),
                        "dt": parser_montant(ligne.get("dt")),
                        "cheque": parser_montant(ligne.get("cheque")),
                        "banque": parser_montant(ligne.get("banque")),
                        "reste": parser_montant(ligne.get("reste")),
                        "operateur": operateur_courant,
                        "page": num_page,
                        "ligne_source": ligne_jointe,
                    })
    except Exception as e:
        log(f"ERREUR lors de la lecture du PDF : {e}")
        return None, None, None, None

    if not en_tete_detecte:
        return None, None, None, None

    return lignes_principales, lignes_rappel, totaux_pdf, periode_controlee


def _auto_controle_totaux(df_analysis, totaux_pdf):
    """Vérifie la cohérence des totaux."""
    if df_analysis.empty or not totaux_pdf or "prime" not in df_analysis.columns:
        return

    try:
        somme_calculee = df_analysis["prime"].dropna().sum()
    except Exception:
        return

    ligne_totaux = next((t for t in totaux_pdf if t.startswith("Totaux")), None)
    if not ligne_totaux:
        return

    nombres = re.findall(r"[\d\s]+,\d{2}", ligne_totaux)
    if not nombres:
        return
    total_pdf = parser_montant(nombres[0])
    if total_pdf is None:
        return

    ecart = round(abs(somme_calculee - total_pdf), 2)
    if ecart <= 0.5:
        log(f"AUTO-CONTRÔLE OK : total extrait ({somme_calculee:.2f}) = total PDF ({total_pdf:.2f}).")
    else:
        log(f"** ATTENTION AUTO-CONTRÔLE ** : écart de {ecart} DH entre extrait et PDF.")


def lire_pdf_encaissements(chemin_pdf):
    log(f"Lecture du PDF d'encaissements : {chemin_pdf}")
    lignes, rappels, totaux_pdf, periode_controlee = extraire_etat_encaissements(chemin_pdf)

    if lignes is None:
        log("ERREUR : impossible de lire le PDF (mise en page non reconnue).")
        return pd.DataFrame(), pd.DataFrame(), None

    df_analysis = pd.DataFrame(lignes)
    df_rappels = pd.DataFrame(rappels) if rappels else pd.DataFrame()
    log(f"{len(df_analysis)} ligne(s) extraite(s), {len(df_rappels)} ligne(s) de rappel.")

    if periode_controlee:
        log(f"Période : du {periode_controlee[0]} au {periode_controlee[1]}")

    _auto_controle_totaux(df_analysis, totaux_pdf)

    return df_analysis, df_rappels, periode_controlee


# ================================================================================
# LECTURE DES RAPPORTS COMPAGNIES (EXCEL OU TEXTE TABULÉ)
# ================================================================================

ALIAS_COLONNES_CIE = {
    "police": ["n police", "n° police", "police"],
    "attestation": ["attestation"],
    "usage": ["usage"],
    # 'souscripteur' en premier : pour MAROC ASSISTANCE, "Assuré principal" vaut
    # "Conducteur" (générique) alors que "Souscripteur" contient le vrai nom.
    "client": ["souscripteur", "client", "assure principal", "assure", "assuré"],
    "prime_prorata": ["p. prorata", "p prorata", "prime prorata", "prorata"],
    "prime_annuelle": ["p. annuelle", "p annuelle", "prime annuelle", "prime ttc", "prime"],
    "date_effet": ["d. effet", "date effet", "date debut", "date début"],
    "confirmee": ["confirmee", "confirmée"],
    "annulee": ["annullee", "annulee", "annullée", "annulée"],
    "controle": ["control"],
}


def _trouver_colonne(en_tetes_normalises, alias_list):
    for alias in alias_list:
        for i, h in enumerate(en_tetes_normalises):
            if alias in h:
                return i
    return None


def _detecter_ligne_entete(df):
    """
    Cherche la ligne qui contient les vrais en-têtes (police, attestation, client...)
    quand le fichier a des lignes de titre au-dessus du tableau.
    Retourne l'index de la ligne d'en-tête, ou None si les colonnes sont déjà bonnes.
    """
    mots_cles = ["police", "attestation", "client", "assure", "prime", "quittance", "effet"]

    cols = " ".join(normaliser_texte(str(c)).lower() for c in df.columns)
    if sum(1 for m in mots_cles if m in cols) >= 2:
        return None  # les colonnes actuelles sont déjà les bons en-têtes

    for i in range(min(25, len(df))):
        ligne = " ".join(normaliser_texte(str(v)).lower() for v in df.iloc[i].tolist())
        if sum(1 for m in mots_cles if m in ligne) >= 2:
            return i
    return None


def _lire_table_brute(chemin_fichier):
    """
    Lit un fichier Excel, CSV ou texte - détection automatique du format,
    du séparateur (tabulation, point-virgule, virgule) et de la ligne d'en-tête.
    """
    chemin_str = str(chemin_fichier)
    brut = None

    # 1) Excel natif (.xlsx / .xls) - lire TOUTES les feuilles sans supposer
    #    l'emplacement de l'en-tête (header=None), garder la plus remplie.
    try:
        feuilles = pd.read_excel(chemin_str, dtype=str, header=None, sheet_name=None)
        meilleure = None
        for _nom_feuille, df_f in feuilles.items():
            if df_f is None or df_f.empty:
                continue
            df_f = df_f.dropna(how="all").dropna(axis=1, how="all")
            if df_f.empty or len(df_f.columns) < 2:
                continue
            score = df_f.notna().sum().sum()  # nombre de cellules remplies
            if meilleure is None or score > meilleure[0]:
                meilleure = (score, df_f)
        if meilleure is not None:
            brut = meilleure[1].reset_index(drop=True)
    except Exception as e:
        log(f"   (lecture pandas Excel échouée: {e} — tentative openpyxl directe)")
        brut = None

    # 1bis) Repli : openpyxl en lecture directe (contourne les soucis de version
    #       pandas/openpyxl, ex. 'Workbook contains no default style').
    if brut is None:
        try:
            import warnings
            from openpyxl import load_workbook
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = load_workbook(chemin_str, read_only=True, data_only=True)
                meilleure = None
                for ws in wb.worksheets:
                    lignes = []
                    for row in ws.iter_rows(values_only=True):
                        lignes.append(["" if v is None else str(v) for v in row])
                    if not lignes:
                        continue
                    largeur = max(len(l) for l in lignes)
                    lignes = [l + [""] * (largeur - len(l)) for l in lignes]
                    df_f = pd.DataFrame(lignes)
                    df_f = df_f.replace("", pd.NA).dropna(how="all").dropna(axis=1, how="all")
                    if df_f.empty or len(df_f.columns) < 2:
                        continue
                    score = df_f.notna().sum().sum()
                    if meilleure is None or score > meilleure[0]:
                        meilleure = (score, df_f)
                wb.close()
                if meilleure is not None:
                    brut = meilleure[1].fillna("").reset_index(drop=True)
        except Exception as e:
            log(f"   (lecture openpyxl directe échouée: {e})")
            brut = None

    # 2) HTML déguisé en Excel (fréquent avec les exports des compagnies)
    if brut is None:
        try:
            tables = pd.read_html(chemin_str)
            tables = [t for t in tables if len(t.columns) >= 3]
            if tables:
                brut = max(tables, key=len)
                brut = brut.where(brut.notna(), "").astype(str)
        except Exception:
            brut = None

    # 3) CSV / texte : détecter le séparateur et la ligne d'en-tête
    if brut is None:
        contenu = None
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                with open(chemin_str, "r", encoding=enc, errors="ignore") as f:
                    contenu = f.read()
                break
            except Exception:
                continue

        if contenu:
            lignes = contenu.splitlines()
            meilleur = None
            for sep in ("\t", ";", ","):
                # Trouver la première ligne avec assez de séparateurs (l'en-tête)
                idx = None
                for i, l in enumerate(lignes[:50]):
                    if l.count(sep) >= 4:
                        idx = i
                        break
                if idx is None:
                    continue
                try:
                    df = pd.read_csv(io.StringIO("\n".join(lignes[idx:])),
                                     dtype=str, sep=sep, engine="python",
                                     on_bad_lines="skip")
                    if len(df.columns) >= 4 and (meilleur is None or len(df.columns) > len(meilleur.columns)):
                        meilleur = df
                except Exception:
                    continue
            brut = meilleur

    if brut is None or brut.empty:
        return None

    # Si les vrais en-têtes sont plus bas (lignes de titre au-dessus du tableau)
    ligne_entete = _detecter_ligne_entete(brut)
    if ligne_entete is not None:
        nouvelles = [str(v) if v is not None else "" for v in brut.iloc[ligne_entete].tolist()]
        brut = brut.iloc[ligne_entete + 1:].reset_index(drop=True)
        brut.columns = nouvelles

    # Supprimer les lignes entièrement vides
    brut = brut.dropna(how="all")
    return brut if not brut.empty else None


def lire_rapport_compagnie(chemin_fichier, nom_compagnie):
    if chemin_fichier is None:
        return pd.DataFrame()

    log(f"Lecture du rapport {nom_compagnie} : {chemin_fichier}")
    try:
        brut = _lire_table_brute(chemin_fichier)
    except Exception as e:
        log(f"ERREUR pour {nom_compagnie} : {e}")
        return pd.DataFrame()

    if brut is None or brut.empty:
        log(f"ERREUR : impossible de lire un tableau pour {nom_compagnie}.")
        return pd.DataFrame()

    log(f"Colonnes détectées pour {nom_compagnie} : {[str(c)[:25] for c in list(brut.columns)[:12]]}")

    en_tetes_normalises = [normaliser_texte(c).lower() for c in brut.columns]
    idx = {champ: _trouver_colonne(en_tetes_normalises, alias) for champ, alias in ALIAS_COLONNES_CIE.items()}

    def colonne(champ):
        position = idx.get(champ)
        return brut.iloc[:, position] if position is not None else pd.Series([""] * len(brut))

    try:
        standard = pd.DataFrame()
        standard["police"] = colonne("police")

        # Attestation MATU : le numéro (ex 204570168) est complété par la
        # PREMIÈRE LETTRE de la colonne "Usage" (AXX->A, FAX->F, D12->D)
        # pour obtenir l'attestation réelle "A 204570168" (= nom du scan).
        att_brut = list(colonne("attestation"))
        usage_col = list(colonne("usage"))
        att_finale = []
        for i in range(len(att_brut)):
            num = str(att_brut[i] or "").strip()
            usg = str(usage_col[i] if i < len(usage_col) else "" or "").strip()
            if not num or num.lower() == "nan":
                att_finale.append("")
            elif num[:1].isalpha():        # déjà préfixée
                att_finale.append(num)
            elif usg and usg[:1].isalpha():  # ajouter la lettre d'usage
                att_finale.append(f"{usg[:1].upper()} {num}")
            else:
                att_finale.append(num)
        standard["attestation"] = att_finale

        standard["client"] = colonne("client")
        prime_prorata = colonne("prime_prorata").apply(parser_montant)
        prime_annuelle = colonne("prime_annuelle").apply(parser_montant)
        standard["prime"] = prime_prorata.where(prime_prorata.notna(), prime_annuelle)
        standard["prime_annuelle"] = prime_annuelle
        standard["date_effet"] = colonne("date_effet").apply(parser_date)
        standard["confirmee"] = colonne("confirmee")
        standard["annulee"] = colonne("annulee")
        standard["deja_controle"] = colonne("controle")
        standard["compagnie"] = nom_compagnie

        log(f"{len(standard)} ligne(s) lue(s) pour {nom_compagnie}.")
        return standard
    except Exception as e:
        log(f"ERREUR en traitant les colonnes de {nom_compagnie} : {e}")
        return pd.DataFrame()


# ================================================================================
# MOTEUR DE RAPPROCHEMENT
# ================================================================================

def rapprocher(df_analysis, scans_jour, scans_tous, df_compagnies, tolerance_prime, seuil_nom):
    resultats = []

    # Vérifier que les colonnes existent
    if df_analysis.empty or not all(col in df_analysis.columns for col in ["attestation", "police", "assure", "prime"]):
        log("ERREUR : colonnes manquantes dans df_analysis")
        return pd.DataFrame(), pd.DataFrame()

    for _, ligne in df_analysis.iterrows():
        try:
            # Ligne de paiement groupé (chèque global) : simple règlement, pas un contrat
            if str(ligne.get("type_ligne", "contrat")) == "paiement_groupe":
                resultats.append({
                    "N° Quittance (interne)": ligne.get("quittance", ""),
                    "N° Attestation": "",
                    "Assuré": str(ligne.get("assure", "")),
                    "Prime Analysis": ligne.get("prime"),
                    "Chèque": ligne.get("cheque"),
                    "Statut global": "PAIEMENT GROUPÉ",
                    "Détail des problèmes": "Chèque global réglant plusieurs quittances (détaillées en contrats)",
                })
                continue

            cle_attestation = normaliser_cle(ligne.get("attestation", ""))
            cle_police = normaliser_cle(ligne.get("police", ""))
            nom_client = str(ligne.get("assure", ""))
            prime_analysis = ligne.get("prime")
            reste = ligne.get("reste")

            cle_scan_trouvee = None
            if cle_attestation and cle_attestation in scans_tous:
                cle_scan_trouvee = cle_attestation
            elif cle_police and cle_police in scans_tous:
                cle_scan_trouvee = cle_police

            statut_scan = "Oui" if cle_scan_trouvee else "Non"
            nom_fichier_scan = scans_tous.get(cle_scan_trouvee, "")
            scan_manquant = not cle_scan_trouvee
            scan_date_du_jour = cle_scan_trouvee in scans_jour if cle_scan_trouvee else False

            if reste is None or pd.isna(reste):
                statut_encaissement = "INCONNU"
            elif abs(reste) <= tolerance_prime:
                statut_encaissement = "OK"
            else:
                statut_encaissement = "IMPAYE" if reste > 0 else "SURPAYE"

            compagnie_trouvee = None
            ligne_cie = None

            if not df_compagnies.empty:
                # 1) Correspondance par numéro d'attestation
                if cle_attestation and "attestation" in df_compagnies.columns:
                    correspondances = df_compagnies[df_compagnies["attestation"].apply(normaliser_cle) == cle_attestation]
                    if len(correspondances) > 0:
                        ligne_cie = correspondances.iloc[0]
                        compagnie_trouvee = ligne_cie.get("compagnie")

                # 2) Correspondance par numéro de police
                if ligne_cie is None and cle_police and "police" in df_compagnies.columns:
                    correspondances = df_compagnies[df_compagnies["police"].apply(normaliser_cle) == cle_police]
                    if len(correspondances) > 0:
                        ligne_cie = correspondances.iloc[0]
                        compagnie_trouvee = ligne_cie.get("compagnie")

                # 3) Correspondance par nom de client (approximative)
                if ligne_cie is None and nom_client and "client" in df_compagnies.columns:
                    meilleur_score = 0
                    for _, candidate in df_compagnies.iterrows():
                        score_nom = similarite(nom_client, candidate.get("client", ""))
                        if score_nom < seuil_nom:
                            continue
                        prime_cie = candidate.get("prime")
                        if prime_analysis is not None and prime_cie is not None:
                            if abs(prime_analysis - prime_cie) > max(tolerance_prime, prime_analysis * 0.02):
                                continue
                        if score_nom > meilleur_score:
                            meilleur_score = score_nom
                            ligne_cie = candidate
                            compagnie_trouvee = candidate.get("compagnie")

            if ligne_cie is None:
                statut_cie = "NON TROUVE"
                ecart_prime = None
                ecart_nom = None
                statut_police = None
                cie_annulee = ""
            else:
                statut_cie = "TROUVE"
                prime_cie = ligne_cie.get("prime")
                ecart_prime = (round(abs(prime_analysis - prime_cie), 2)
                               if (prime_analysis is not None and prime_cie is not None) else None)
                ecart_nom = round(100 - similarite(nom_client, ligne_cie.get("client", "")), 1)
                statut_police = (cle_police == normaliser_cle(ligne_cie.get("police", "")))
                cie_annulee = str(ligne_cie.get("annulee", ""))

            # FLOTTE / CHÈQUE GROUPÉ : comparer la SOMME des primes (prime_attendue)
            # au total encaissé, au lieu d'une seule ligne compagnie.
            prime_attendue = ligne.get("prime_attendue")
            if prime_attendue is not None and not pd.isna(prime_attendue):
                statut_cie = "TROUVE"
                compagnie_trouvee = compagnie_trouvee or "MATU"
                ecart_prime = (round(abs(prime_analysis - prime_attendue), 2)
                               if prime_analysis is not None else None)
                if ecart_nom is None:
                    ecart_nom = 0.0
                if statut_police is None:
                    statut_police = True

            problemes = []
            if scan_manquant:
                problemes.append("Scan manquant")
            if statut_encaissement == "IMPAYE":
                problemes.append("Impayé")
            if statut_encaissement == "SURPAYE":
                problemes.append("Surpaiement")
            if statut_cie == "NON TROUVE":
                problemes.append("Contrat non retrouvé chez une compagnie")
            if ecart_prime is not None and ecart_prime > tolerance_prime:
                problemes.append(f"Écart de prime ({ecart_prime} DH)")
            if ecart_nom is not None and ecart_nom > (100 - seuil_nom):
                problemes.append("Nom client différent")
            if normaliser_texte(cie_annulee) in ("VRAI", "TRUE", "OUI", "1", "ANNULE"):
                problemes.append("Contrat annulé chez la compagnie")

            if not problemes:
                statut_global = "CONFORME"
            elif any(p in problemes for p in ("Scan manquant", "Contrat non retrouvé chez une compagnie",
                                               "Impayé", "Contrat annulé chez la compagnie")):
                statut_global = "CRITIQUE"
            else:
                statut_global = "ATTENTION"

            resultats.append({
                "N° Quittance (interne)": ligne.get("quittance", ""),
                "N° Quittance (compagnie)": ligne.get("quittance_compagnie", ""),
                "N° Attestation": ligne.get("attestation", ""),
                "Assuré": nom_client,
                "N° Police": ligne.get("police", ""),
                "Prime Analysis": prime_analysis,
                "Espèces": ligne.get("especes"),
                "Chèque": ligne.get("cheque"),
                "Banque": ligne.get("banque"),
                "Reste": reste,
                "Opérateur": ligne.get("operateur", ""),
                "Date effet": ligne.get("date_effet"),
                "Scan présent": statut_scan,
                "Scan daté du jour ?": ("Oui" if scan_date_du_jour else ("Oui, autre date" if not scan_manquant else "")),
                "Nom fichier scan": nom_fichier_scan,
                "Statut encaissement": statut_encaissement,
                "Compagnie retrouvée": compagnie_trouvee or "",
                "Statut compagnie": statut_cie,
                "Écart prime (DH)": ecart_prime,
                "Écart nom (%)": ecart_nom,
                "N° Police identique ?": statut_police,
                "Statut global": statut_global,
                "Détail des problèmes": "; ".join(problemes) if problemes else "Aucun",
            })
        except Exception as e:
            log(f"ERREUR en traitant une ligne : {e}")
            continue

    df_resultat = pd.DataFrame(resultats)

    # Détection des omissions
    omissions = []
    if not df_compagnies.empty and not df_analysis.empty and "police" in df_compagnies.columns and "police" in df_analysis.columns:
        try:
            polices_analysis = set(df_analysis["police"].apply(normaliser_cle))
            for _, cie_ligne in df_compagnies.iterrows():
                cle_p = normaliser_cle(cie_ligne.get("police", ""))
                if cle_p and cle_p not in polices_analysis:
                    omissions.append({
                        "Compagnie": cie_ligne.get("compagnie", ""),
                        "N° Police": cie_ligne.get("police", ""),
                        "N° Attestation": cie_ligne.get("attestation", ""),
                        "Client": cie_ligne.get("client", ""),
                        "Prime (prorata)": cie_ligne.get("prime", ""),
                        "Prime annuelle": cie_ligne.get("prime_annuelle", ""),
                        "Confirmée": cie_ligne.get("confirmee", ""),
                        "Annulée": cie_ligne.get("annulee", ""),
                    })
        except Exception as e:
            log(f"ERREUR en détectant les omissions : {e}")

    df_omissions = pd.DataFrame(omissions)

    return df_resultat, df_omissions


# ================================================================================
# GÉNÉRATION DU RAPPORT EXCEL
# ================================================================================

def _construire_nom_fichier_sortie(periode_controlee):
    horodatage_generation = datetime.now().strftime("%Y-%m-%d_%Hh%M")
    if periode_controlee:
        date_controlee = periode_controlee[0].replace("/", "-")
        return f"Rapprochement_{date_controlee}_genere_{horodatage_generation}.xlsx"
    return f"Rapprochement_DATE-INCONNUE_genere_{horodatage_generation}.xlsx"


COULEUR_CONFORME = "C6E0B4"
COULEUR_ATTENTION = "FFE699"
COULEUR_CRITIQUE = "F4B183"
COULEUR_ENTETE = "1F4E78"


def _style_entete(cellule):
    cellule.fill = PatternFill(start_color=COULEUR_ENTETE, end_color=COULEUR_ENTETE, fill_type="solid")
    cellule.font = Font(color="FFFFFF", bold=True)
    cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ecrire_dataframe(feuille, df, colonne_statut=None, ligne_depart=1):
    if df is None or df.empty:
        feuille.append(["Aucune donnée"])
        return

    # En-têtes
    ligne_entete = ligne_depart
    for j, col in enumerate(df.columns, start=1):
        c = feuille.cell(row=ligne_entete, column=j, value=str(col))
        _style_entete(c)

    # Lignes de données
    for i, (_, ligne) in enumerate(df.iterrows(), start=ligne_entete + 1):
        statut = ligne.get(colonne_statut) if colonne_statut else None
        couleur = {"CONFORME": COULEUR_CONFORME, "ATTENTION": COULEUR_ATTENTION,
                   "CRITIQUE": COULEUR_CRITIQUE}.get(statut)

        for j, col in enumerate(df.columns, start=1):
            valeur = ligne[col]
            try:
                est_vide = pd.isna(valeur)
            except (TypeError, ValueError):
                est_vide = False
            if est_vide:
                valeur = None
            elif isinstance(valeur, datetime):
                valeur = valeur.strftime("%d/%m/%Y")
            c = feuille.cell(row=i, column=j, value=valeur)
            if couleur:
                c.fill = PatternFill(start_color=couleur, end_color=couleur, fill_type="solid")

    # Largeurs colonnes
    for j, col in enumerate(df.columns, start=1):
        largeur = max(12, min(40, len(str(col)) + 4))
        feuille.column_dimensions[get_column_letter(j)].width = largeur

    # Gel des en-têtes
    feuille.freeze_panes = feuille.cell(row=ligne_entete + 1, column=1).coordinate


def generer_rapport_excel(df_resultat, df_omissions, df_rappels, chemin_sortie, periode_controlee=None, df_controle_att=None):
    classeur = Workbook()

    texte_periode = (
        f"Contrôle du {periode_controlee[0]} au {periode_controlee[1]}"
        if periode_controlee else
        "Période contrôlée : NON DÉTECTÉE"
    )

    # Feuille "Rapprochement"
    feuille_rappro = classeur.active
    feuille_rappro.title = "Rapprochement"
    feuille_rappro.append([texte_periode])
    feuille_rappro["A1"].font = Font(bold=True, size=13, color="1F4E78")
    feuille_rappro.append([])
    _ecrire_dataframe(feuille_rappro, df_resultat, colonne_statut="Statut global", ligne_depart=3)

    # Feuille "Écarts"
    feuille_ecarts = classeur.create_sheet("Ecarts")
    feuille_ecarts.append([texte_periode])
    feuille_ecarts["A1"].font = Font(bold=True, size=13, color="1F4E78")
    feuille_ecarts.append([])
    df_ecarts = (df_resultat[~df_resultat["Statut global"].isin(["CONFORME", "PAIEMENT GROUPÉ"])].copy()
                 if (not df_resultat.empty and "Statut global" in df_resultat.columns)
                 else pd.DataFrame())
    _ecrire_dataframe(feuille_ecarts, df_ecarts, colonne_statut="Statut global", ligne_depart=3)

    # Feuille "Omissions"
    feuille_omissions = classeur.create_sheet("Omissions_possibles")
    _ecrire_dataframe(feuille_omissions, df_omissions)

    # Feuille "Quittances Rappel"
    feuille_rappels = classeur.create_sheet("Quittances_Rappel")
    _ecrire_dataframe(feuille_rappels, df_rappels)

    # Feuille "Synthèse"
    feuille_synthese = classeur.create_sheet("Synthese")
    total = len(df_resultat) if not df_resultat.empty else 0

    # Calcul des statistiques SÉCURISÉ
    if not df_resultat.empty and "Statut global" in df_resultat.columns:
        nb_conforme = int((df_resultat["Statut global"] == "CONFORME").sum())
        nb_attention = int((df_resultat["Statut global"] == "ATTENTION").sum())
        nb_critique = int((df_resultat["Statut global"] == "CRITIQUE").sum())
    else:
        nb_conforme = nb_attention = nb_critique = 0

    taux_conformite = round(100 * nb_conforme / total, 1) if total > 0 else 0

    lignes_synthese = [
        ["RAPPORT DE CONTRÔLE JOURNALIER", ""],
        ["Date/période contrôlée", f"du {periode_controlee[0]} au {periode_controlee[1]}" if periode_controlee else "NON DÉTECTÉE"],
        ["Date de génération", datetime.now().strftime("%d/%m/%Y %H:%M")],
        [],
        ["Indicateur", "Valeur"],
        ["Nombre total de lignes", total],
        ["Conformes", nb_conforme],
        ["Attention", nb_attention],
        ["Critiques", nb_critique],
        ["Taux de conformité (%)", taux_conformite],
        ["Omissions possibles", len(df_omissions) if not df_omissions.empty else 0],
        ["Quittances rappel", len(df_rappels) if not df_rappels.empty else 0],
    ]

    # Statistiques du contrôle des attestations scannées
    if df_controle_att is not None and not df_controle_att.empty:
        att_ok = int((df_controle_att["Statut"] == "✅ TROUVÉE").sum())
        att_manq = int((df_controle_att["Statut"] == "❌ MANQUANTE").sum())
        total_att = len(df_controle_att)
        lignes_synthese += [
            [],
            ["CONTRÔLE DES SCANS", ""],
            ["Attestations scannées (trouvées)", att_ok],
            ["Attestations NON scannées (manquantes)", att_manq],
            ["Taux de scan (%)", round(100 * att_ok / total_att, 1) if total_att else 0],
        ]

    for ligne in lignes_synthese:
        feuille_synthese.append(ligne)
    feuille_synthese.column_dimensions["A"].width = 60
    feuille_synthese.column_dimensions["B"].width = 25
    for cellule in feuille_synthese["A1:B1"][0]:
        cellule.font = Font(bold=True, size=14)
    for cellule in feuille_synthese["A5:B5"][0]:
        _style_entete(cellule)

    # Feuille "Contrôle Attestations" (NOUVEAU)
    if df_controle_att is not None and not df_controle_att.empty:
        feuille_att = classeur.create_sheet("Controle_Attestations", 1)
        feuille_att.append(["CONTRÔLE DES ATTESTATIONS SCANNÉES"])
        feuille_att["A1"].font = Font(bold=True, size=13, color="C00000")
        feuille_att.append(["Vérification que tous les documents ont été scannés"])
        feuille_att.append([])

        # Ajouter les données de contrôle
        _ecrire_dataframe(feuille_att, df_controle_att, colonne_statut="Statut", ligne_depart=4)

        # Ajouter un résumé
        att_trouvees = (df_controle_att["Statut"] == "✅ TROUVÉE").sum()
        att_manquantes = (df_controle_att["Statut"] == "❌ MANQUANTE").sum()
        total_att = len(df_controle_att)

        feuille_att.append([])
        feuille_att.append(["RÉSUMÉ"])
        feuille_att.append(["Attestations trouvées (scannées)", att_trouvees])
        feuille_att.append(["Attestations manquantes (NON scannées)", att_manquantes])
        feuille_att.append(["Total", total_att])
        feuille_att.append(["Taux de conformité (%)", round(100 * att_trouvees / total_att, 1) if total_att > 0 else 0])

        feuille_att.column_dimensions["A"].width = 20
        feuille_att.column_dimensions["B"].width = 18
        feuille_att.column_dimensions["C"].width = 28
        feuille_att.column_dimensions["D"].width = 14
        feuille_att.column_dimensions["E"].width = 14
        feuille_att.column_dimensions["F"].width = 24

    classeur.save(chemin_sortie)
    log(f"Rapport Excel créé : {chemin_sortie}")


# ================================================================================
# CONTRÔLE DES ATTESTATIONS SCANNÉES
# ================================================================================

def controler_attestations_scannees(df_analysis, df_compagnies, attestations_dispo=None, df_rappels=None):
    r"""
    Contrôle si toutes les attestations ont été scannées.
    Compare les numéros d'attestation/police avec les fichiers dans \\KARIMA\images analisis

    - Lignes principales (MATU, SANLAM...) : scan nommé par n° d'attestation
    - Quittances encaissée / MAROC ASSISTANCE : scan nommé par n° de police
      (format IAL.xx.xxxxxx). Ces lignes n'ont pas de n° d'attestation.

    Args:
        attestations_dispo: dict {numero: Path} déjà chargé (évite un 2e scan réseau)
        df_rappels: quittances encaissée (section MAROC ASSISTANCE du PDF)

    Returns:
        DataFrame de contrôle
    """
    log("\n" + "=" * 70)
    log("CONTRÔLE DES ATTESTATIONS SCANNÉES")
    log("=" * 70)

    # Récupérer les attestations depuis config_sources (sauf si déjà fournies)
    if attestations_dispo is None:
        try:
            attestations_dispo = SourcesData.chercher_attestations() or {}
        except Exception as e:
            log(f"❌ Erreur lors de la recherche d'attestations : {e}")
            attestations_dispo = {}
    log(f"✅ {len(attestations_dispo)} attestation(s) trouvée(s) sur le réseau")

    controle = []
    vus = set()
    nb_trouvees = 0
    nb_manquantes = 0

    def traiter_ligne(att_brut, pol_brut, assure):
        nonlocal nb_trouvees, nb_manquantes
        att_brut = str(att_brut or "").strip()
        pol_brut = str(pol_brut or "").strip()
        cle_att = normaliser_cle(att_brut)
        cle_pol = normaliser_cle(pol_brut)

        cle_unique = cle_att or cle_pol
        if not cle_unique or cle_unique in vus:
            return
        vus.add(cle_unique)

        # Chercher le scan : par attestation d'abord, puis par police (MAROC)
        chemin_trouve = None
        cle_trouvee = None
        if cle_att and cle_att in attestations_dispo:
            chemin_trouve = attestations_dispo[cle_att]
            cle_trouvee = "attestation"
        elif cle_pol and cle_pol in attestations_dispo:
            chemin_trouve = attestations_dispo[cle_pol]
            cle_trouvee = "police"

        trouvee = chemin_trouve is not None

        controle.append({
            "Numéro Attestation": att_brut if att_brut else "(sans attestation)",
            "Numéro Police": pol_brut,
            "Assuré": str(assure or ""),
            "Statut": "✅ TROUVÉE" if trouvee else "❌ MANQUANTE",
            "Trouvé par": cle_trouvee or "---",
            "Fichier": chemin_trouve.name if trouvee else "---",
        })

        if trouvee:
            nb_trouvees += 1
        else:
            nb_manquantes += 1

    # Table police -> attestation depuis les rapports compagnies (surtout MATU).
    # Pour les chèques groupés, l'attestation n'est PAS dans l'encaissement :
    # on la récupère ici via le numéro de police.
    corr_police_att = {}
    if df_compagnies is not None and not df_compagnies.empty \
            and "police" in df_compagnies.columns and "attestation" in df_compagnies.columns:
        for _, l in df_compagnies.iterrows():
            p = normaliser_cle(str(l.get("police", "")))
            a = str(l.get("attestation", "") or "").strip()
            if p and a and p not in corr_police_att:
                corr_police_att[p] = a

    # 1) Lignes principales (attestation directe de l'encaissement)
    #    On saute les lignes de paiement groupé (chèque global, pas un document à scanner).
    if not df_analysis.empty and "attestation" in df_analysis.columns:
        for _, ligne in df_analysis.iterrows():
            if str(ligne.get("type_ligne", "contrat")) == "paiement_groupe":
                continue
            # Flotte : vérifier CHAQUE attestation (un scan par véhicule)
            atts_flotte = ligne.get("attestations_flotte") if "attestations_flotte" in ligne.index else None
            if isinstance(atts_flotte, (list, tuple)) and len(atts_flotte) > 0:
                for att in atts_flotte:
                    traiter_ligne(att, ligne.get("police", ""), ligne.get("assure", ""))
            else:
                traiter_ligne(ligne.get("attestation", ""), ligne.get("police", ""), ligne.get("assure", ""))

    # 2) Sous-quittances (chèques groupés + Quittances encaissée)
    if df_rappels is not None and not df_rappels.empty and "police" in df_rappels.columns:
        for _, ligne in df_rappels.iterrows():
            police = str(ligne.get("police", "") or "").strip()
            quittance = str(ligne.get("quittance", "") or "").strip()

            # Pour MAROC : police = attestation = quittance sans le suffixe "-1"
            if not police and quittance:
                police = re.sub(r"-\d+$", "", quittance)

            # Résoudre l'attestation :
            #  - MATU/autres : depuis le rapport compagnie (par police)
            #  - MAROC ASSISTANCE : la police EST l'attestation (format non numérique, ex IAL.xx)
            att = corr_police_att.get(normaliser_cle(police), "")
            if not att and police and not police[:1].isdigit():
                att = police
            traiter_ligne(att, police, ligne.get("assure", "") or "(quittance de rappel)")

    if not controle:
        log("Aucune ligne d'encaissement à contrôler.")
        return pd.DataFrame()

    total = nb_trouvees + nb_manquantes

    # Résumé du contrôle
    log(f"\n📊 RÉSUMÉ DU CONTRÔLE D'ATTESTATIONS :")
    log(f"   ✅ Trouvées : {nb_trouvees}/{total}")
    log(f"   ❌ Manquantes : {nb_manquantes}/{total}")

    if nb_manquantes > 0:
        log(f"\n⚠️  ALERTES - {nb_manquantes} document(s) non scanné(s) :")
        for item in controle:
            if "MANQUANTE" in item["Statut"]:
                ref = item["Numéro Attestation"]
                if ref == "(sans attestation)" and item["Numéro Police"]:
                    ref = f"police {item['Numéro Police']}"
                elif item["Numéro Police"]:
                    ref += f" (police {item['Numéro Police']})"
                log(f"   ❌ {ref} - {item['Assuré']}")

    return pd.DataFrame(controle) if controle else pd.DataFrame()


def enrichir_rappels_avec_assure(df_rappels, df_compagnies):
    """
    Ajoute le nom de l'assuré aux quittances de rappel en cherchant la police
    dans les rapports compagnies (surtout MAROC ASSISTANCE, colonne 'client').
    """
    if df_rappels is None or df_rappels.empty:
        return df_rappels
    if "assure" not in df_rappels.columns:
        df_rappels = df_rappels.copy()
        df_rappels["assure"] = ""

    if df_compagnies is None or df_compagnies.empty or "police" not in df_compagnies.columns:
        return df_rappels

    # Table de correspondance police normalisée -> nom client
    corr = {}
    for _, l in df_compagnies.iterrows():
        cle = normaliser_cle(str(l.get("police", "")))
        nom = str(l.get("client", "") or "").strip()
        if cle and nom and cle not in corr:
            corr[cle] = nom

    if not corr:
        return df_rappels

    df_rappels = df_rappels.copy()
    nb = 0
    for i, l in df_rappels.iterrows():
        if str(l.get("assure", "") or "").strip():
            continue
        cle = normaliser_cle(str(l.get("police", "")))
        if cle in corr:
            df_rappels.at[i, "assure"] = corr[cle]
            nb += 1
    if nb:
        log(f"{nb} nom(s) d'assuré ajouté(s) aux quittances de rappel.")
    return df_rappels


def integrer_cheques_groupes(df_analysis, df_rappels, df_compagnies):
    """
    OPTION 1 : les sous-quittances d'un chèque groupé deviennent de VRAIS
    contrats à rapprocher, et la ligne de paiement 'CAR' (chèque global)
    est neutralisée (marquée PAIEMENT, plus signalée comme anomalie).

    - MATU/autres : l'attestation de la sous-quittance est reprise du rapport
      compagnie (via la police).
    - MAROC ASSISTANCE : police = attestation = quittance sans le '-1'.

    Retourne df_analysis enrichi (avec colonne 'type_ligne').
    """
    if df_analysis is None or df_analysis.empty:
        return df_analysis

    df_analysis = df_analysis.copy()
    if "type_ligne" not in df_analysis.columns:
        df_analysis["type_ligne"] = "contrat"

    # Marquer les lignes de paiement groupé : assuré finissant par 'CAR' + police vide
    def _est_paiement(l):
        ass = normaliser_texte(str(l.get("assure", "")))
        pol = normaliser_cle(str(l.get("police", "")))
        return ass.endswith("CAR") and not pol
    masque = df_analysis.apply(_est_paiement, axis=1)
    nb_paie = int(masque.sum())
    df_analysis.loc[masque, "type_ligne"] = "paiement_groupe"

    # Table police -> {attestations (avec lettre), primes, client} depuis les compagnies.
    # Une flotte a PLUSIEURS lignes (véhicules) sous la même police.
    corr = {}
    if df_compagnies is not None and not df_compagnies.empty and "police" in df_compagnies.columns:
        for _, l in df_compagnies.iterrows():
            p = normaliser_cle(str(l.get("police", "")))
            if not p:
                continue
            info = corr.setdefault(p, {"atts": [], "primes": [], "client": ""})
            att = str(l.get("attestation", "") or "").strip()
            if att and att.lower() != "nan":
                info["atts"].append(att)
            prime = l.get("prime")
            if prime is not None and not pd.isna(prime):
                info["primes"].append(float(prime))
            if not info["client"]:
                info["client"] = str(l.get("client", "") or "").strip()

    # Regrouper les sous-quittances par police (une flotte = une police, plusieurs quittances)
    groupes = defaultdict(list)
    if df_rappels is not None and not df_rappels.empty:
        for _, l in df_rappels.iterrows():
            police = str(l.get("police", "") or "").strip()
            quittance = str(l.get("quittance", "") or "").strip()
            if not police and quittance:
                police = re.sub(r"-\d+$", "", quittance)
            groupes[police].append(l)

    nouvelles = []
    for police, lignes in groupes.items():
        cle_pol = normaliser_cle(police)
        info = corr.get(cle_pol)

        # Total encaissé = somme des montants des sous-quittances de la flotte
        total_paye = 0.0
        quittances = []
        for l in lignes:
            m = l.get("montant_encaisse")
            if m is None or pd.isna(m):
                m = l.get("prime")
            if m is not None and not pd.isna(m):
                total_paye += float(m)
            q = str(l.get("quittance", "") or "").strip()
            if q:
                quittances.append(q)

        # Somme des primes MATU de la flotte (règle : somme des primes vs encaissement)
        prime_attendue = round(sum(info["primes"]), 2) if info and info["primes"] else None
        atts = list(info["atts"]) if info else []
        client = info["client"] if info else ""

        # Attestation principale (pour l'affichage) et liste pour le contrôle des scans
        if atts:
            att_principale = atts[0]
        elif police and not police[:1].isdigit():
            att_principale = police   # MAROC : police = attestation
            atts = [police]
        else:
            att_principale = ""

        assure = ""
        for l in lignes:
            a = str(l.get("assure", "") or "").strip()
            if a:
                assure = a
                break
        assure = assure or client

        nouvelles.append({
            "quittance": quittances[0] if quittances else "",
            "quittance_compagnie": " + ".join(quittances),
            "attestation": att_principale,
            "attestations_flotte": atts,          # pour le contrôle des scans
            "assure": assure,
            "police": police,
            "date_effet": lignes[0].get("date_effet"),
            "prime": round(total_paye, 2),
            "prime_attendue": prime_attendue,     # somme des primes MATU de la flotte
            "reference_banque": "",
            "especes": 0.0,
            "dt": 0.0,
            "cheque": round(total_paye, 2),
            "banque": 0.0,
            "reste": 0.0,
            "operateur": "",
            "page": None,
            "ligne_source": f"Chèque groupé / flotte ({len(lignes)} quittance(s), police {police})",
            "type_ligne": "contrat",
        })

    if nouvelles:
        df_analysis = pd.concat([df_analysis, pd.DataFrame(nouvelles)], ignore_index=True)

    log(f"Chèques groupés : {nb_paie} ligne(s) de paiement neutralisée(s), "
        f"{len(nouvelles)} contrat(s)/flotte(s) intégré(s) depuis {sum(len(v) for v in groupes.values())} sous-quittance(s).")
    return df_analysis


# ================================================================================
# PROGRAMME PRINCIPAL
# ================================================================================

def main():
    log("=" * 70)
    log("DÉMARRAGE DU RAPPROCHEMENT AUTOMATIQUE (VERSION CORRIGÉE)")
    log("=" * 70)

    config = charger_config("config.ini")

    output_path = config["Chemins"].get("output_path", "./Rapports_Generes")
    os.makedirs(output_path, exist_ok=True)

    tolerance_prime = float(config["Tolerances"].get("tolerance_prime", "2.0"))
    seuil_nom = float(config["Tolerances"].get("seuil_similarite_nom", "80"))

    date_cible = demander_date_controle(config["Controle"].get("date_controle", "auto"))
    filtrer_scans_par_date = config["Controle"].get("filtrer_scans_par_date", "oui").strip().lower() in ("oui", "yes", "true", "1")

    # Convertir la date en format DD.MM.YYYY pour les sources
    date_search = None
    if date_cible and SOURCES_DISPONIBLES:
        date_str = date_cible.strftime("%d.%m.%Y") if hasattr(date_cible, 'strftime') else str(date_cible)
        date_search = date_str.replace("/", ".") if "/" in date_str else date_str

    df_analysis = pd.DataFrame()
    df_rappels = pd.DataFrame()
    periode_controlee = None
    attestations_reseau = None

    # NOUVELLE APPROCHE: utiliser config_sources si disponible
    if SOURCES_DISPONIBLES:
        log("✅ Utilisation des sources configurées (Google Drive + Réseau)")

        # Chercher encaissements
        chemin_pdf = SourcesData.chercher_encaissements(date_search)
        if chemin_pdf:
            df_analysis, df_rappels, periode_controlee = lire_pdf_encaissements(str(chemin_pdf))
        else:
            log("ARRÊT PARTIEL : aucun PDF d'encaissements trouvé.")

        # Les scans = attestations trouvées sur le réseau \\KARIMA\images analisis
        scans_jour, scans_tous = {}, {}
        try:
            attestations_reseau = SourcesData.chercher_attestations() or {}
            scans_tous = {cle: chemin.name for cle, chemin in attestations_reseau.items()}
            log(f"{len(scans_tous)} attestation(s) scannée(s) disponibles sur le réseau.")

            # Déterminer les scans datés du jour contrôlé (via la date du fichier)
            if date_cible is not None and not df_analysis.empty and "attestation" in df_analysis.columns:
                for att in df_analysis["attestation"].dropna():
                    cle = normaliser_cle(str(att))
                    chemin_att = attestations_reseau.get(cle)
                    if chemin_att:
                        try:
                            date_fichier = datetime.fromtimestamp(chemin_att.stat().st_mtime).date()
                            if date_fichier == date_cible:
                                scans_jour[cle] = chemin_att.name
                        except Exception:
                            pass
        except Exception as e:
            log(f"ATTENTION : impossible de lire les attestations réseau : {e}")
            attestations_reseau = None

        # Chercher rapports compagnies
        rapports = SourcesData.chercher_rapports_compagnies(date_search)

        morceaux_cies = []
        for nom_cie, chemin_cie in rapports.items():
            if chemin_cie:
                df_cie = lire_rapport_compagnie(str(chemin_cie), nom_cie)
                if not df_cie.empty:
                    morceaux_cies.append(df_cie)
    else:
        # ANCIENNE APPROCHE: utiliser les chemins de config
        log("⚠️ config_sources.py non disponible, utilisation des chemins config")
        drive_path = config["Chemins"].get("drive_path", ".")
        scans_path = config["Chemins"].get("scans_path", ".")

        motif_pdf = config["Fichiers"].get(
            "pattern_encaissements",
            "Etat des encaissements*.pdf;Etat_des_encaissements*.pdf;Rapport encaissement*.pdf",
        )
        chemin_pdf, date_pdf_confirmee = trouver_fichier_pour_date(drive_path, motif_pdf, date_cible)

        if chemin_pdf:
            df_analysis, df_rappels, periode_controlee = lire_pdf_encaissements(chemin_pdf)
        else:
            log("ARRÊT PARTIEL : aucun PDF d'encaissements trouvé.")

        scans_jour, scans_tous = lister_fichiers_scans(scans_path, date_cible, filtrer_scans_par_date)

        # Lecture des compagnies
        compagnies = {
            "MATU": config["Fichiers"].get("pattern_matu", "RAPPORT MATU*.xlsx"),
            "SANLAM": config["Fichiers"].get("pattern_sanlam", "RAPPORT SANLAM*.xlsx"),
            "WAFA ASSURANCE": config["Fichiers"].get("pattern_wafa", "RAPPORT WAFA*.xlsx"),
            "MAROC ASSISTANCE": config["Fichiers"].get("pattern_maroc_assistance", "RAPPORT MAROC ASSISTANCE*.xlsx"),
        }

        morceaux_cies = []
        for nom_cie, motif in compagnies.items():
            chemin_cie, _ = trouver_fichier_pour_date(drive_path, motif, date_cible)
            df_cie = lire_rapport_compagnie(chemin_cie, nom_cie)
            if not df_cie.empty:
                morceaux_cies.append(df_cie)

    df_compagnies = pd.concat(morceaux_cies, ignore_index=True) if morceaux_cies else pd.DataFrame()

    # Enrichir les quittances de rappel avec le nom d'assuré (via le rapport compagnie, par police)
    df_rappels = enrichir_rappels_avec_assure(df_rappels, df_compagnies)

    # OPTION 1 : intégrer les sous-quittances des chèques groupés comme vrais
    # contrats, et neutraliser les lignes de paiement 'CAR'.
    df_analysis = integrer_cheques_groupes(df_analysis, df_rappels, df_compagnies)

    # Rapprochement
    if df_analysis.empty:
        log("Impossible de poursuivre : aucune donnée Analysis.")
        df_resultat, df_omissions = pd.DataFrame(), pd.DataFrame()
    else:
        df_resultat, df_omissions = rapprocher(df_analysis, scans_jour, scans_tous, df_compagnies, tolerance_prime, seuil_nom)

    # Contrôle des scans : les sous-quittances sont désormais dans df_analysis,
    # on ne repasse donc PAS df_rappels (éviter les doublons).
    df_controle_att = controler_attestations_scannees(df_analysis, df_compagnies, attestations_reseau, None)

    # Génération du rapport Excel
    nom_fichier_sortie = _construire_nom_fichier_sortie(periode_controlee)
    chemin_sortie = os.path.join(output_path, nom_fichier_sortie)
    generer_rapport_excel(df_resultat, df_omissions, df_rappels, chemin_sortie, periode_controlee, df_controle_att)

    ecrire_journal(output_path)

    log("=" * 70)
    log("TERMINÉ.")
    log(f"Ouvrez le fichier : {chemin_sortie}")
    log("=" * 70)

    try:
        os.startfile(chemin_sortie)
    except Exception:
        pass

    return chemin_sortie


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print("\n" + "=" * 70)
        print("UNE ERREUR EST SURVENUE :")
        print("=" * 70)
        traceback.print_exc()
        print("\nCopiez ce message et envoyez-le pour assistance.")
    try:
        input("\nAppuyez sur Entrée pour fermer...")
    except EOFError:
        pass
